# -*- coding: utf-8 -*-
"""Загрузка одного документа или набора раздельных документов как корпуса.

Публикация разбита на отдельные документы по разделам
(src/split_publication.py, каталог sections/), поэтому проверки должны
работать либо с одним файлом, либо со всем каталогом: абзацы, таблицы и
рисунки всех документов склеиваются в один корпус.
"""
from pathlib import Path

from docx import Document


class Corpus:
    """Единый корпус: один DOCX или все DOCX каталога.

    Предоставляет те же свойства, что и Document (paragraphs, tables,
    inline_shapes), так что проверочные скрипты не замечают подмены.
    """

    def __init__(self, path):
        self.path = Path(path)
        if self.path.is_dir():
            self.paths = sorted(self.path.glob('*.docx'))
        else:
            self.paths = [self.path]
        if not self.paths:
            raise SystemExit('не найдено ни одного DOCX: %s' % path)
        self.docs = [Document(p) for p in self.paths]

    @property
    def paragraphs(self):
        out = []
        for d in self.docs:
            out.extend(d.paragraphs)
        return out

    @property
    def tables(self):
        out = []
        for d in self.docs:
            out.extend(d.tables)
        return out

    @property
    def inline_shapes(self):
        out = []
        for d in self.docs:
            out.extend(d.inline_shapes)
        return out
